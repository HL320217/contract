"""
数据源处理模块 - 处理Excel、CSV、JSON等数据源
"""

pd = None
import math
import json
import csv
import os
from typing import Dict, List, Any, Optional
import logging
from config import SUPPORTED_DATA_FORMATS, MAX_DATA_SIZE, ERROR_MESSAGES

# 配置日志
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class DataProcessor:
    """数据源处理器"""
    
    def __init__(self):
        self.data = None
        self.columns = []
        self.row_count = 0
        self.last_error = ""
        self._pandas_mode = True
        
    def load_data(self, file_path: str) -> bool:
        """
        加载数据源文件
        
        Args:
            file_path: 数据源文件路径
            
        Returns:
            bool: 加载是否成功
        """
        try:
            global pd
            if pd is None:
                try:
                    import pandas as _pd
                    pd = _pd
                    self._pandas_mode = True
                except Exception as e:
                    pd = None
                    self._pandas_mode = False
                    self.last_error = f"pandas不可用，将使用兼容模式: {str(e)}"
            # 检查文件是否存在
            if not os.path.exists(file_path):
                self.last_error = f"文件不存在: {file_path}"
                logger.error(self.last_error)
                return False
            
            # 检查文件大小
            file_size = os.path.getsize(file_path) / (1024 * 1024)  # MB
            if file_size > MAX_DATA_SIZE:
                self.last_error = f"文件大小超出限制: {file_size:.2f}MB > {MAX_DATA_SIZE}MB"
                logger.error(self.last_error)
                return False
            
            # 获取文件扩展名
            _, ext = os.path.splitext(file_path.lower())
            
            if ext not in SUPPORTED_DATA_FORMATS:
                self.last_error = f"不支持的文件格式: {ext}"
                logger.error(self.last_error)
                return False
            
            # 根据文件类型加载数据
            if ext == '.xlsx':
                if self._pandas_mode:
                    try:
                        self.data = pd.read_excel(file_path, engine='openpyxl')
                    except Exception:
                        self.data = pd.read_excel(file_path)
                else:
                    try:
                        from openpyxl import load_workbook
                        wb = load_workbook(file_path, read_only=True, data_only=True)
                        ws = wb.active
                        rows = list(ws.iter_rows(values_only=True))
                        if not rows:
                            self.last_error = "Excel文件为空"
                            return False
                        headers = ["" if h is None else str(h) for h in rows[0]]
                        records = []
                        for row in rows[1:]:
                            rec = {}
                            for i, v in enumerate(row):
                                rec[headers[i]] = "" if v is None else str(v)
                            records.append(rec)
                        self.data = records
                    except Exception as e:
                        self.last_error = f"读取xlsx失败: {str(e)}"
                        return False
            elif ext == '.xls':
                if self._pandas_mode:
                    try:
                        self.data = pd.read_excel(file_path, engine='xlrd')
                    except Exception:
                        self.data = pd.read_excel(file_path)
                else:
                    try:
                        import xlrd
                        book = xlrd.open_workbook(file_path)
                        sheet = book.sheet_by_index(0)
                        headers = [str(sheet.cell_value(0, c)) for c in range(sheet.ncols)]
                        records = []
                        for r in range(1, sheet.nrows):
                            rec = {}
                            for c in range(sheet.ncols):
                                v = sheet.cell_value(r, c)
                                rec[headers[c]] = "" if v is None else str(v)
                            records.append(rec)
                        self.data = records
                    except Exception as e:
                        self.last_error = f"读取xls失败: {str(e)}"
                        return False
            elif ext == '.csv':
                encodings = ['utf-8', 'utf-8-sig', 'gbk', 'gb2312']
                if self._pandas_mode:
                    loaded = False
                    for encoding in encodings:
                        try:
                            self.data = pd.read_csv(
                                file_path,
                                encoding=encoding,
                                encoding_errors='ignore',
                                engine='python',
                                dtype=str,
                                on_bad_lines='skip'
                            )
                            loaded = True
                            break
                        except TypeError:
                            try:
                                self.data = pd.read_csv(file_path, encoding=encoding)
                                loaded = True
                                break
                            except Exception:
                                continue
                        except Exception:
                            continue
                    if not loaded:
                        try:
                            with open(file_path, 'r', encoding='utf-8', errors='ignore', newline='') as f:
                                reader = csv.reader(f)
                                rows = list(reader)
                                if not rows:
                                    self.last_error = "CSV为空或不可读取"
                                    return False
                                headers = rows[0]
                                records = [dict(zip(headers, r)) for r in rows[1:]]
                                self.data = pd.DataFrame(records)
                        except Exception as e:
                            self.last_error = f"CSV读取失败: {str(e)}"
                            return False
                else:
                    try:
                        for enc in encodings:
                            try:
                                with open(file_path, 'r', encoding=enc, errors='ignore', newline='') as f:
                                    reader = csv.reader(f)
                                    rows = list(reader)
                                    if rows:
                                        headers = rows[0]
                                        records = [dict(zip(headers, r)) for r in rows[1:]]
                                        self.data = records
                                        break
                            except Exception:
                                continue
                        if self.data is None:
                            self.last_error = "CSV读取失败"
                            return False
                    except Exception as e:
                        self.last_error = f"CSV读取失败: {str(e)}"
                        return False
            elif ext == '.json':
                with open(file_path, 'r', encoding='utf-8') as f:
                    json_data = json.load(f)
                    if self._pandas_mode:
                        if isinstance(json_data, list):
                            self.data = pd.DataFrame(json_data)
                        elif isinstance(json_data, dict):
                            self.data = pd.DataFrame([json_data])
                        else:
                            self.last_error = "JSON格式不正确"
                            return False
                    else:
                        if isinstance(json_data, list):
                            self.data = [
                                {k: ("" if v is None else str(v)) for k, v in row.items()}
                                for row in json_data
                            ]
                        elif isinstance(json_data, dict):
                            self.data = [{k: ("" if v is None else str(v)) for k, v in json_data.items()}]
                        else:
                            self.last_error = "JSON格式不正确"
                            return False
            
            # 更新数据信息
            if self._pandas_mode:
                self.columns = list(self.data.columns)
                self.row_count = len(self.data)
            else:
                self.columns = list(self.data[0].keys()) if self.data else []
                self.row_count = len(self.data)
            
            logger.info(f"成功加载数据: {self.row_count}行, {len(self.columns)}列")
            return True
            
        except Exception as e:
            self.last_error = f"加载数据时发生错误: {str(e)}"
            logger.error(self.last_error)
            return False
    
    def get_data_preview(self, max_rows: int = 5) -> Dict[str, Any]:
        """
        获取数据预览
        
        Args:
            max_rows: 最大预览行数
            
        Returns:
            Dict: 包含预览数据的字典
        """
        if self.data is None:
            return {"error": "未加载数据"}
        
        try:
            if self._pandas_mode:
                preview_data = self.data.head(max_rows)
                return {
                    "columns": self.columns,
                    "row_count": self.row_count,
                    "preview": preview_data.to_dict('records'),
                    "sample_row": self.data.iloc[0].to_dict() if self.row_count > 0 else {}
                }
            else:
                return {
                    "columns": self.columns,
                    "row_count": self.row_count,
                    "preview": self.data[:max_rows],
                    "sample_row": self.data[0] if self.row_count > 0 else {}
                }
        except Exception as e:
            logger.error(f"获取数据预览时发生错误: {str(e)}")
            return {"error": str(e)}
    
    def get_row_data(self, row_index: int) -> Optional[Dict[str, Any]]:
        """
        获取指定行的数据
        
        Args:
            row_index: 行索引
            
        Returns:
            Dict: 行数据字典，如果出错返回None
        """
        if self.data is None:
            return None
        
        try:
            if 0 <= row_index < self.row_count:
                if self._pandas_mode:
                    row_data = self.data.iloc[row_index].to_dict()
                    cleaned = {}
                    for k, v in row_data.items():
                        if v is None or (isinstance(v, float) and math.isnan(v)):
                            cleaned[k] = ""
                        else:
                            cleaned[k] = str(v)
                    return cleaned
                else:
                    row = self.data[row_index]
                    cleaned = {k: ("" if v is None else str(v)) for k, v in row.items()}
                    return cleaned
            else:
                logger.error(f"行索引超出范围: {row_index}")
                return None
        except Exception as e:
            logger.error(f"获取行数据时发生错误: {str(e)}")
            return None
    
    def get_all_data(self) -> List[Dict[str, Any]]:
        """
        获取所有数据
        
        Returns:
            List[Dict]: 所有行数据的列表
        """
        if self.data is None:
            return []
        
        try:
            all_data = []
            for i in range(self.row_count):
                row_data = self.get_row_data(i)
                if row_data:
                    all_data.append(row_data)
            return all_data
        except Exception as e:
            logger.error(f"获取所有数据时发生错误: {str(e)}")
            return []
    
    def validate_columns(self, required_columns: List[str]) -> Dict[str, Any]:
        """
        验证数据是否包含必需的列
        
        Args:
            required_columns: 必需的列名列表
            
        Returns:
            Dict: 验证结果
        """
        if self.data is None:
            return {"valid": False, "error": "未加载数据"}
        
        missing_columns = []
        for col in required_columns:
            if col not in self.columns:
                missing_columns.append(col)
        
        if missing_columns:
            return {
                "valid": False,
                "error": f"缺少必需的列: {', '.join(missing_columns)}",
                "missing_columns": missing_columns,
                "available_columns": self.columns
            }
        
        return {"valid": True, "message": "所有必需列都存在"}
    
    def get_column_mapping_suggestions(self, template_variables: List[str]) -> Dict[str, str]:
        """
        为模板变量建议列映射
        
        Args:
            template_variables: 模板中的变量列表
            
        Returns:
            Dict: 变量到列的映射建议
        """
        if self.data is None:
            return {}
        
        mapping = {}
        
        for var in template_variables:
            # 尝试精确匹配
            if var in self.columns:
                mapping[var] = var
                continue
            
            # 尝试模糊匹配
            var_lower = var.lower()
            for col in self.columns:
                col_lower = col.lower()
                if var_lower in col_lower or col_lower in var_lower:
                    mapping[var] = col
                    break
        
        return mapping
